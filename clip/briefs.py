"""Read campaign briefs (PDF / Word / text / Google Doc) and pull out structured rules."""
import re
import urllib.request
from pathlib import Path

from . import config, llm

MAX_CHARS = 60_000
GDOC = re.compile(r"^https://docs\.google\.com/document/d/([\w-]+)")


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in t.rows)
    else:
        text = path.read_text(errors="replace")
    return text.strip()[:MAX_CHARS]


def fetch_gdoc(url: str) -> str:
    """Link-shared Google Docs only. The URL is rebuilt from the doc id, never fetched as given."""
    m = GDOC.match(url.strip())
    if not m:
        raise ValueError("Not a Google Docs link (expected https://docs.google.com/document/d/...)")
    export = f"https://docs.google.com/document/d/{m.group(1)}/export?format=txt"
    req = urllib.request.Request(export, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=15) as r:
        if "text/plain" not in r.headers.get("Content-Type", ""):
            raise ValueError("Couldn't read that doc. Set sharing to 'Anyone with the link', or download it as PDF/Word and drop the file.")
        return r.read(2_000_000).decode("utf-8", "replace").strip()[:MAX_CHARS]


TOOL = {
    "name": "set_brief",
    "description": "Structured campaign requirements extracted from the brief.",
    "input_schema": {"type": "object", "properties": {
        "rules": {"type": "string", "description": "Concise bullet list of everything a clipper must do or avoid: required mentions/tags, banned content, style, length, platforms, how clips are judged, payout rules"},
        "hashtags": {"type": "array", "items": {"type": "string"}},
        "disclosure": {"type": "string", "description": "Required disclosure tag such as #ad, or empty"},
        "platform": {"type": "string", "enum": ["TikTok", "Instagram", "YouTube Shorts", "Any"]},
        "rate_cpm": {"type": "number", "description": "Dollars per 1000 views, 0 if unknown"},
        "min_seconds": {"type": "integer"}, "max_seconds": {"type": "integer"},
    }, "required": ["rules"]},
}


def analyze(text: str) -> dict:
    """Return suggested campaign fields. Without an API key, just hand back the raw text."""
    if not text:
        return {}
    tags = list(dict.fromkeys(re.findall(r"#\w+", text)))[:12]
    if not config.ANTHROPIC_API_KEY:
        return {"rules": text[:3000], "hashtags": tags, "note": "No API key: raw text used."}
    import anthropic
    out, _ = llm.call_tool(
        anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY), config.CLAUDE_MODEL, TOOL,
        "Extract the requirements of this clipping campaign brief. Only include what the "
        "brief states; omit fields it doesn't mention.\n\n<brief>\n" + text + "\n</brief>", 1500)
    return {k: v for k, v in out.items() if v not in (None, "", [], 0)}
